"""
Utilidades para generar reportes en PDF y Excel
"""
from django.http import HttpResponse
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime


def generate_sales_pdf(sales, filename='reporte_ventas.pdf'):
    """
    Genera un reporte de ventas en formato PDF
    """
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    doc = SimpleDocTemplate(response, pagesize=letter)
    elements = []
    styles = getSampleStyleSheet()

    # Título
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        alignment=TA_CENTER,
        spaceAfter=30
    )
    title = Paragraph('Reporte de Ventas - RubiPerfumeria', title_style)
    elements.append(title)

    # Fecha de generación
    date_style = ParagraphStyle(
        'DateStyle',
        parent=styles['Normal'],
        alignment=TA_RIGHT,
        fontSize=10
    )
    date_text = Paragraph(
        f'Generado: {datetime.now().strftime("%d/%m/%Y %H:%M")}',
        date_style
    )
    elements.append(date_text)
    elements.append(Spacer(1, 0.3 * inch))

    # Resumen
    total_sales = sales.count()
    total_amount = sum(sale.total for sale in sales)

    summary_data = [
        ['Resumen del Reporte'],
        [f'Total de ventas: {total_sales}'],
        [f'Monto total: ${total_amount:.2f}']
    ]
    summary_table = Table(summary_data, colWidths=[6 * inch])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 12),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('GRID', (0, 0), (-1, -1), 1, colors.black)
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 0.3 * inch))

    # Tabla de ventas
    if sales:
        data = [['Ticket', 'Fecha/Hora', 'Cajero', 'Método', 'Total', 'Descuento', 'Estado']]

        for sale in sales[:100]:  # Limitar a 100 ventas por reporte
            data.append([
                sale.numero_ticket,
                sale.fecha_hora.strftime('%d/%m/%Y %H:%M'),
                sale.cajero.username,
                sale.get_metodo_pago_display(),
                f'${sale.total:.2f}',
                f'${sale.descuento:.2f}',
                sale.get_estado_display()
            ])

        table = Table(data, colWidths=[1*inch, 1.3*inch, 1*inch, 1.2*inch, 1*inch, 1*inch, 1*inch])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('FONTSIZE', (0, 1), (-1, -1), 8)
        ]))
        elements.append(table)

    doc.build(elements)
    return response


def generate_sales_excel(sales, filename='reporte_ventas.xlsx'):
    """
    Genera un reporte de ventas en formato Excel
    """
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    wb = Workbook()
    ws = wb.active
    ws.title = 'Reporte de Ventas'

    # Estilos
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')

    # Título
    ws['A1'] = 'Reporte de Ventas - RubiPerfumeria'
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells('A1:G1')

    # Fecha
    ws['A2'] = f'Generado: {datetime.now().strftime("%d/%m/%Y %H:%M")}'
    ws.merge_cells('A2:G2')

    # Resumen
    total_sales = sales.count()
    total_amount = sum(sale.total for sale in sales)
    ws['A4'] = 'Resumen'
    ws['A4'].font = Font(bold=True)
    ws['A5'] = f'Total de ventas: {total_sales}'
    ws['A6'] = f'Monto total: ${total_amount:.2f}'

    # Encabezados
    headers = ['Ticket', 'Fecha y Hora', 'Cajero', 'Método de Pago', 'Total', 'Descuento', 'Estado']
    ws.append([])  # Línea en blanco
    ws.append(headers)

    header_row = ws.max_row
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # Datos
    for sale in sales:
        ws.append([
            sale.numero_ticket,
            sale.fecha_hora.strftime('%d/%m/%Y %H:%M'),
            sale.cajero.username,
            sale.get_metodo_pago_display(),
            float(sale.total),
            float(sale.descuento),
            sale.get_estado_display()
        ])

    # Ajustar ancho de columnas
    column_widths = [15, 18, 15, 18, 12, 12, 15]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width

    wb.save(response)
    return response


def generate_inventory_pdf(products, filename='reporte_inventario.pdf'):
    """
    Genera un reporte de inventario en formato PDF
    """
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    doc = SimpleDocTemplate(response, pagesize=letter)
    elements = []
    styles = getSampleStyleSheet()

    # Título
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        alignment=TA_CENTER,
        spaceAfter=30
    )
    title = Paragraph('Reporte de Inventario - RubiPerfumeria', title_style)
    elements.append(title)

    # Fecha de generación
    date_style = ParagraphStyle(
        'DateStyle',
        parent=styles['Normal'],
        alignment=TA_RIGHT,
        fontSize=10
    )
    date_text = Paragraph(
        f'Generado: {datetime.now().strftime("%d/%m/%Y %H:%M")}',
        date_style
    )
    elements.append(date_text)
    elements.append(Spacer(1, 0.3 * inch))

    # Resumen
    total_products = products.count()
    low_stock_count = sum(1 for p in products if p.requiere_reabastecimiento)
    total_value = sum(p.valor_inventario for p in products)

    summary_data = [
        ['Resumen del Inventario'],
        [f'Total de productos: {total_products}'],
        [f'Productos con bajo stock: {low_stock_count}'],
        [f'Valor total del inventario: ${total_value:.2f}']
    ]
    summary_table = Table(summary_data, colWidths=[6 * inch])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 12),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('GRID', (0, 0), (-1, -1), 1, colors.black)
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 0.3 * inch))

    # Tabla de productos
    if products:
        data = [['Código', 'Producto', 'Stock', 'Stock Mín.', 'Precio', 'Valor', 'Estado']]

        for product in products[:100]:  # Limitar a 100 productos
            estado = 'BAJO STOCK' if product.requiere_reabastecimiento else (
                'Activo' if product.activo else 'Inactivo'
            )
            data.append([
                product.codigo_barras[:10],
                product.nombre[:20],
                str(product.stock_actual),
                str(product.stock_minimo),
                f'${product.precio:.2f}',
                f'${product.valor_inventario:.2f}',
                estado
            ])

        table = Table(data, colWidths=[1*inch, 1.5*inch, 0.8*inch, 0.8*inch, 1*inch, 1*inch, 1.2*inch])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('FONTSIZE', (0, 1), (-1, -1), 7)
        ]))
        elements.append(table)

    doc.build(elements)
    return response


def generate_inventory_excel(products, filename='reporte_inventario.xlsx'):
    """
    Genera un reporte de inventario en formato Excel
    """
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    wb = Workbook()
    ws = wb.active
    ws.title = 'Reporte de Inventario'

    # Estilos
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')

    # Título
    ws['A1'] = 'Reporte de Inventario - RubiPerfumeria'
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells('A1:I1')

    # Fecha
    ws['A2'] = f'Generado: {datetime.now().strftime("%d/%m/%Y %H:%M")}'
    ws.merge_cells('A2:I2')

    # Resumen
    total_products = products.count()
    low_stock_count = sum(1 for p in products if p.requiere_reabastecimiento)
    total_value = sum(p.valor_inventario for p in products)

    ws['A4'] = 'Resumen'
    ws['A4'].font = Font(bold=True)
    ws['A5'] = f'Total de productos: {total_products}'
    ws['A6'] = f'Productos con bajo stock: {low_stock_count}'
    ws['A7'] = f'Valor total del inventario: ${total_value:.2f}'

    # Encabezados
    headers = [
        'Código de Barras', 'Producto', 'Marca', 'Categoría',
        'Stock Actual', 'Stock Mínimo', 'Precio', 'Valor Total', 'Estado'
    ]
    ws.append([])  # Línea en blanco
    ws.append(headers)

    header_row = ws.max_row
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # Datos
    for product in products:
        estado = 'BAJO STOCK' if product.requiere_reabastecimiento else (
            'Activo' if product.activo else 'Inactivo'
        )
        ws.append([
            product.codigo_barras,
            product.nombre,
            product.marca,
            product.get_categoria_display(),
            product.stock_actual,
            product.stock_minimo,
            float(product.precio),
            float(product.valor_inventario),
            estado
        ])

    # Ajustar ancho de columnas
    column_widths = [18, 25, 15, 15, 12, 12, 12, 12, 15]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width

    wb.save(response)
    return response
