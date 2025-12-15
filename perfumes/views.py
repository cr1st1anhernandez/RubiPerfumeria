from django.shortcuts import render, get_object_or_404, redirect
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from accounts.decorators import supervisor_required
from .models import Perfume
from .forms import PerfumeForm


@supervisor_required
def perfume_list(request):
    """Lista de perfumes con opción de exportar inventario."""
    exportar = request.GET.get('exportar')

    perfumes = Perfume.objects.all().order_by('nombre')

    # Obtener movimientos de inventario
    from pos.models import MovimientoInventario
    movimientos = MovimientoInventario.objects.select_related(
        'perfume', 'usuario', 'venta'
    ).order_by('-fecha_creacion')

    # Si se solicita exportación
    if exportar == 'pdf':
        return exportar_inventario_pdf(perfumes, movimientos)
    elif exportar == 'excel':
        return exportar_inventario_excel(perfumes, movimientos)

    context = {
        'perfumes': perfumes,
        'movimientos': movimientos[:50],  # Últimos 50 movimientos
        'total_movimientos': movimientos.count(),
    }

    return render(request, 'perfumes/perfume_list.html', context)


def exportar_inventario_pdf(perfumes, movimientos):
    """Exporta el inventario a PDF."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from io import BytesIO
    from django.utils import timezone

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    elements = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#2c3e50'),
        spaceAfter=30,
        alignment=TA_CENTER
    )

    title = Paragraph("Reporte de Inventario - Rubi Perfumería", title_style)
    elements.append(title)

    fecha_texto = f"Generado: {timezone.now().strftime('%d/%m/%Y %H:%M')}"
    fecha = Paragraph(fecha_texto, styles['Normal'])
    elements.append(fecha)
    elements.append(Spacer(1, 20))

    # Stock Actual
    subtitle = Paragraph("Stock Actual de Productos", styles['Heading2'])
    elements.append(subtitle)
    elements.append(Spacer(1, 12))

    stock_data = [['Producto', 'Marca', 'Stock', 'Precio', 'Valor Total']]
    total_valor_inventario = 0

    for producto in perfumes:
        valor_total = float(producto.precio) * producto.stock
        total_valor_inventario += valor_total
        stock_data.append([
            producto.nombre,
            producto.marca,
            str(producto.stock),
            f'${producto.precio:,.2f}',
            f'${valor_total:,.2f}'
        ])

    # Agregar fila de total
    stock_data.append(['', '', '', 'TOTAL:', f'${total_valor_inventario:,.2f}'])

    stock_table = Table(stock_data, colWidths=[2*inch, 1.5*inch, 0.8*inch, 1.2*inch, 1.2*inch])
    stock_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3498db')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('ALIGN', (2, 0), (2, -1), 'CENTER'),
        ('ALIGN', (3, 0), (-1, -1), 'RIGHT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -2), colors.beige),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#ecf0f1')),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 1, colors.white),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
    ]))

    elements.append(stock_table)
    elements.append(PageBreak())

    # Movimientos
    subtitle2 = Paragraph("Últimos Movimientos de Inventario", styles['Heading2'])
    elements.append(subtitle2)
    elements.append(Spacer(1, 12))

    mov_data = [['Fecha', 'Producto', 'Tipo', 'Cantidad', 'Usuario']]
    for mov in movimientos[:50]:
        mov_data.append([
            mov.fecha_creacion.strftime('%d/%m/%Y %H:%M'),
            f"{mov.perfume.marca} {mov.perfume.nombre}",
            mov.get_tipo_movimiento_display(),
            str(mov.cantidad),
            mov.usuario.username
        ])

    mov_table = Table(mov_data, colWidths=[1.5*inch, 2*inch, 1.2*inch, 1*inch, 1*inch])
    mov_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3498db')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('ALIGN', (3, 0), (3, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.white),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
    ]))

    elements.append(mov_table)

    if movimientos.count() > 50:
        nota = Paragraph(f"<i>Mostrando los primeros 50 de {movimientos.count()} movimientos</i>", styles['Normal'])
        elements.append(Spacer(1, 12))
        elements.append(nota)

    doc.build(elements)
    buffer.seek(0)

    response = HttpResponse(buffer, content_type='application/pdf')
    filename = f"inventario_{timezone.now().strftime('%Y%m%d')}.pdf"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    return response


def exportar_inventario_excel(perfumes, movimientos):
    """Exporta el inventario a Excel."""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from io import BytesIO
    from django.utils import timezone

    wb = openpyxl.Workbook()
    ws_stock = wb.active
    ws_stock.title = "Inventario"

    header_font = Font(name='Arial', size=14, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='3498DB', end_color='3498DB', fill_type='solid')
    title_font = Font(name='Arial', size=16, bold=True)

    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Título
    ws_stock['A1'] = 'Inventario de Productos - Rubi Perfumería'
    ws_stock['A1'].font = title_font
    ws_stock['A1'].alignment = Alignment(horizontal='center')
    ws_stock.merge_cells('A1:G1')

    ws_stock['A3'] = 'Fecha:'
    ws_stock['B3'] = timezone.now().strftime('%d/%m/%Y %H:%M')
    ws_stock['A3'].font = Font(bold=True)

    # Encabezados
    headers = ['Producto', 'Marca', 'Tipo', 'Volumen (ml)', 'Stock', 'Precio Unitario', 'Valor Total']
    for col, header in enumerate(headers, start=1):
        cell = ws_stock.cell(row=5, column=col)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = border

    # Datos
    row = 6
    total_valor = 0
    for producto in perfumes:
        valor_total = float(producto.precio) * producto.stock
        total_valor += valor_total

        ws_stock.cell(row=row, column=1).value = producto.nombre
        ws_stock.cell(row=row, column=2).value = producto.marca
        ws_stock.cell(row=row, column=3).value = producto.tipo
        ws_stock.cell(row=row, column=4).value = producto.volumen
        ws_stock.cell(row=row, column=5).value = producto.stock
        ws_stock.cell(row=row, column=6).value = float(producto.precio)
        ws_stock.cell(row=row, column=7).value = valor_total

        ws_stock.cell(row=row, column=6).number_format = '$#,##0.00'
        ws_stock.cell(row=row, column=7).number_format = '$#,##0.00'

        for col in range(1, 8):
            ws_stock.cell(row=row, column=col).border = border

        row += 1

    # Total
    ws_stock.cell(row=row, column=6).value = 'TOTAL:'
    ws_stock.cell(row=row, column=6).font = Font(bold=True)
    ws_stock.cell(row=row, column=6).alignment = Alignment(horizontal='right')
    ws_stock.cell(row=row, column=7).value = total_valor
    ws_stock.cell(row=row, column=7).number_format = '$#,##0.00'
    ws_stock.cell(row=row, column=7).font = Font(bold=True)

    for col in range(6, 8):
        ws_stock.cell(row=row, column=col).border = border

    # Ajustar anchos
    ws_stock.column_dimensions['A'].width = 25
    ws_stock.column_dimensions['B'].width = 20
    ws_stock.column_dimensions['C'].width = 10
    ws_stock.column_dimensions['D'].width = 12
    ws_stock.column_dimensions['E'].width = 10
    ws_stock.column_dimensions['F'].width = 15
    ws_stock.column_dimensions['G'].width = 15

    # Hoja 2: Movimientos
    ws_mov = wb.create_sheet(title="Movimientos")

    ws_mov['A1'] = 'Movimientos de Inventario'
    ws_mov['A1'].font = title_font
    ws_mov['A1'].alignment = Alignment(horizontal='center')
    ws_mov.merge_cells('A1:G1')

    headers_mov = ['Fecha', 'Hora', 'Producto', 'Marca', 'Tipo Movimiento', 'Cantidad', 'Usuario']
    for col, header in enumerate(headers_mov, start=1):
        cell = ws_mov.cell(row=3, column=col)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = border

    row = 4
    for mov in movimientos:
        ws_mov.cell(row=row, column=1).value = mov.fecha_creacion.strftime('%d/%m/%Y')
        ws_mov.cell(row=row, column=2).value = mov.fecha_creacion.strftime('%H:%M:%S')
        ws_mov.cell(row=row, column=3).value = mov.perfume.nombre
        ws_mov.cell(row=row, column=4).value = mov.perfume.marca
        ws_mov.cell(row=row, column=5).value = mov.get_tipo_movimiento_display()
        ws_mov.cell(row=row, column=6).value = mov.cantidad
        ws_mov.cell(row=row, column=7).value = mov.usuario.username

        for col in range(1, 8):
            ws_mov.cell(row=row, column=col).border = border

        row += 1

    ws_mov.column_dimensions['A'].width = 12
    ws_mov.column_dimensions['B'].width = 10
    ws_mov.column_dimensions['C'].width = 25
    ws_mov.column_dimensions['D'].width = 20
    ws_mov.column_dimensions['E'].width = 18
    ws_mov.column_dimensions['F'].width = 10
    ws_mov.column_dimensions['G'].width = 15

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"inventario_{timezone.now().strftime('%Y%m%d')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    return response

    response = HttpResponse(
        buffer,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"inventario_{timezone.now().strftime('%Y%m%d')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    return response


def exportar_inventario_pdf(perfumes, movimientos):
    """Exporta el inventario a PDF."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from io import BytesIO
    from django.utils import timezone

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    elements = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#2c3e50'),
        spaceAfter=30,
        alignment=TA_CENTER
    )

    title = Paragraph("Reporte de Inventario - Rubi Perfumería", title_style)
    elements.append(title)

    fecha_texto = f"Generado: {timezone.now().strftime('%d/%m/%Y %H:%M')}"
    fecha = Paragraph(fecha_texto, styles['Normal'])
    elements.append(fecha)
    elements.append(Spacer(1, 20))

    # Stock Actual
    subtitle = Paragraph("Stock Actual de Productos", styles['Heading2'])
    elements.append(subtitle)
    elements.append(Spacer(1, 12))

    stock_data = [['Producto', 'Marca', 'Stock', 'Precio', 'Valor Total']]
    total_valor_inventario = 0

    for producto in perfumes:
        valor_total = float(producto.precio) * producto.stock
        total_valor_inventario += valor_total
        stock_data.append([
            producto.nombre,
            producto.marca,
            str(producto.stock),
            f'${producto.precio:,.2f}',
            f'${valor_total:,.2f}'
        ])

    # Agregar fila de total
    stock_data.append(['', '', '', 'TOTAL:', f'${total_valor_inventario:,.2f}'])

    stock_table = Table(stock_data, colWidths=[2*inch, 1.5*inch, 0.8*inch, 1.2*inch, 1.2*inch])
    stock_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3498db')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('ALIGN', (2, 0), (2, -1), 'CENTER'),
        ('ALIGN', (3, 0), (-1, -1), 'RIGHT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -2), colors.beige),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#ecf0f1')),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 1, colors.white),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
    ]))

    elements.append(stock_table)
    elements.append(PageBreak())

    # Movimientos
    subtitle2 = Paragraph("Últimos Movimientos de Inventario", styles['Heading2'])
    elements.append(subtitle2)
    elements.append(Spacer(1, 12))

    mov_data = [['Fecha', 'Producto', 'Tipo', 'Cantidad', 'Usuario']]
    for mov in movimientos[:50]:
        mov_data.append([
            mov.fecha_creacion.strftime('%d/%m/%Y %H:%M'),
            f"{mov.perfume.marca} {mov.perfume.nombre}",
            mov.get_tipo_movimiento_display(),
            str(mov.cantidad),
            mov.usuario.username
        ])

    mov_table = Table(mov_data, colWidths=[1.5*inch, 2*inch, 1.2*inch, 1*inch, 1*inch])
    mov_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3498db')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('ALIGN', (3, 0), (3, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.white),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
    ]))

    elements.append(mov_table)

    if movimientos.count() > 50:
        nota = Paragraph(f"<i>Mostrando los primeros 50 de {movimientos.count()} movimientos</i>", styles['Normal'])
        elements.append(Spacer(1, 12))
        elements.append(nota)

    doc.build(elements)
    buffer.seek(0)

    response = HttpResponse(buffer, content_type='application/pdf')
    filename = f"inventario_{timezone.now().strftime('%Y%m%d')}.pdf"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    return response


def exportar_inventario_excel(perfumes, movimientos):
    """Exporta el inventario a Excel."""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from io import BytesIO
    from django.utils import timezone

    wb = openpyxl.Workbook()
    ws_stock = wb.active
    ws_stock.title = "Inventario"

    header_font = Font(name='Arial', size=14, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='3498DB', end_color='3498DB', fill_type='solid')
    title_font = Font(name='Arial', size=16, bold=True)

    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Título
    ws_stock['A1'] = 'Inventario de Productos - Rubi Perfumería'
    ws_stock['A1'].font = title_font
    ws_stock['A1'].alignment = Alignment(horizontal='center')
    ws_stock.merge_cells('A1:G1')

    ws_stock['A3'] = 'Fecha:'
    ws_stock['B3'] = timezone.now().strftime('%d/%m/%Y %H:%M')
    ws_stock['A3'].font = Font(bold=True)

    # Encabezados
    headers = ['Producto', 'Marca', 'Tipo', 'Volumen (ml)', 'Stock', 'Precio Unitario', 'Valor Total']
    for col, header in enumerate(headers, start=1):
        cell = ws_stock.cell(row=5, column=col)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = border

    # Datos
    row = 6
    total_valor = 0
    for producto in perfumes:
        valor_total = float(producto.precio) * producto.stock
        total_valor += valor_total

        ws_stock.cell(row=row, column=1).value = producto.nombre
        ws_stock.cell(row=row, column=2).value = producto.marca
        ws_stock.cell(row=row, column=3).value = producto.tipo
        ws_stock.cell(row=row, column=4).value = producto.volumen
        ws_stock.cell(row=row, column=5).value = producto.stock
        ws_stock.cell(row=row, column=6).value = float(producto.precio)
        ws_stock.cell(row=row, column=7).value = valor_total

        ws_stock.cell(row=row, column=6).number_format = '$#,##0.00'
        ws_stock.cell(row=row, column=7).number_format = '$#,##0.00'

        for col in range(1, 8):
            ws_stock.cell(row=row, column=col).border = border

        row += 1

    # Total
    ws_stock.cell(row=row, column=6).value = 'TOTAL:'
    ws_stock.cell(row=row, column=6).font = Font(bold=True)
    ws_stock.cell(row=row, column=6).alignment = Alignment(horizontal='right')
    ws_stock.cell(row=row, column=7).value = total_valor
    ws_stock.cell(row=row, column=7).number_format = '$#,##0.00'
    ws_stock.cell(row=row, column=7).font = Font(bold=True)

    for col in range(6, 8):
        ws_stock.cell(row=row, column=col).border = border

    # Ajustar anchos
    ws_stock.column_dimensions['A'].width = 25
    ws_stock.column_dimensions['B'].width = 20
    ws_stock.column_dimensions['C'].width = 10
    ws_stock.column_dimensions['D'].width = 12
    ws_stock.column_dimensions['E'].width = 10
    ws_stock.column_dimensions['F'].width = 15
    ws_stock.column_dimensions['G'].width = 15

    # Hoja 2: Movimientos
    ws_mov = wb.create_sheet(title="Movimientos")

    ws_mov['A1'] = 'Movimientos de Inventario'
    ws_mov['A1'].font = title_font
    ws_mov['A1'].alignment = Alignment(horizontal='center')
    ws_mov.merge_cells('A1:G1')

    headers_mov = ['Fecha', 'Hora', 'Producto', 'Marca', 'Tipo Movimiento', 'Cantidad', 'Usuario']
    for col, header in enumerate(headers_mov, start=1):
        cell = ws_mov.cell(row=3, column=col)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = border

    row = 4
    for mov in movimientos:
        ws_mov.cell(row=row, column=1).value = mov.fecha_creacion.strftime('%d/%m/%Y')
        ws_mov.cell(row=row, column=2).value = mov.fecha_creacion.strftime('%H:%M:%S')
        ws_mov.cell(row=row, column=3).value = mov.perfume.nombre
        ws_mov.cell(row=row, column=4).value = mov.perfume.marca
        ws_mov.cell(row=row, column=5).value = mov.get_tipo_movimiento_display()
        ws_mov.cell(row=row, column=6).value = mov.cantidad
        ws_mov.cell(row=row, column=7).value = mov.usuario.username

        for col in range(1, 8):
            ws_mov.cell(row=row, column=col).border = border

        row += 1

    ws_mov.column_dimensions['A'].width = 12
    ws_mov.column_dimensions['B'].width = 10
    ws_mov.column_dimensions['C'].width = 25
    ws_mov.column_dimensions['D'].width = 20
    ws_mov.column_dimensions['E'].width = 18
    ws_mov.column_dimensions['F'].width = 10
    ws_mov.column_dimensions['G'].width = 15

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"inventario_{timezone.now().strftime('%Y%m%d')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    return response


@supervisor_required
def perfume_detail(request, pk):
    perfume = get_object_or_404(Perfume, pk=pk)
    return render(request, 'perfumes/perfume_detail.html', {'perfume': perfume})


@supervisor_required
def perfume_create(request):
    if request.method == 'POST':
        form = PerfumeForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Perfume creado exitosamente.')
            return redirect('perfume_list')
    else:
        form = PerfumeForm()
    return render(request, 'perfumes/perfume_form.html', {'form': form, 'action': 'Crear'})


@supervisor_required
def perfume_update(request, pk):
    perfume = get_object_or_404(Perfume, pk=pk)
    if request.method == 'POST':
        form = PerfumeForm(request.POST, instance=perfume)
        if form.is_valid():
            form.save()
            messages.success(request, 'Perfume actualizado exitosamente.')
            return redirect('perfume_detail', pk=pk)
    else:
        form = PerfumeForm(instance=perfume)
    return render(request, 'perfumes/perfume_form.html', {'form': form, 'action': 'Editar', 'perfume': perfume})


@supervisor_required
def perfume_delete(request, pk):
    perfume = get_object_or_404(Perfume, pk=pk)
    if request.method == 'POST':
        perfume.delete()
        messages.success(request, 'Perfume eliminado exitosamente.')
        return redirect('perfume_list')
    return render(request, 'perfumes/perfume_confirm_delete.html', {'perfume': perfume})
